#!/usr/bin/env python3
"""Build report/initial_proposal.pptx (4:3, styled after last semester's sample deck).

Slides 5-12 (Proposed Trading Simulation, owner Hao Wang) are generated from
report/proposal_trading_simulation.md. Edit that file, not this script, to change them:
  ## heading     -> one slide          > line      -> takeaway under the title
  ### heading    -> card on the slide  - bullet    -> on the slide
  $$ ... $$      -> equation image     | table |   -> native table
  indented bullets and plain paragraphs -> speaker notes
The other team sections are condensed from the team's Google-Doc draft (TEAM_* below).

Needs python-pptx, matplotlib and Pillow, plus Poppler's `pdfimages` for the paper figure:
    python3 report/build_initial_proposal.py
"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from PIL import Image, ImageChops  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "report" / "proposal_trading_simulation.md"
PAPER_PDF = ROOT / "Pairs_Trading_RV_Arbitrage_Rule.pdf"
OUT = ROOT / "report" / "initial_proposal.pptx"

# Palette taken from sample_initial_proposal.pptx; Arial throughout so LibreOffice previews match PowerPoint widths.
BLUE, NAVY, INK, GREEN = "2A75B8", "0B5394", "1C1917", "047857"
MUTED, TINT, TINT2, GREEN_TINT, WHITE = "5F6B7A", "EEF3F9", "D6E4F3", "E3F2EC", "FFFFFF"
TITLE_FONT, BODY_FONT = "Arial", "Arial"
FOOTER = "IEOR 4576 · Extending Gatev, Goetzmann & Rouwenhorst (2006)"
HAO = "Hao Wang"

plt.rcParams["mathtext.fontset"] = "cm"
TMP = Path(tempfile.mkdtemp(prefix="proposal_deck_"))


# --------------------------------------------------------------------------- markdown source


@dataclass
class Block:
    title: str
    takeaway: str = ""
    bullets: list[str] = field(default_factory=list)
    eqs: list[str] = field(default_factory=list)
    table: list[list[str]] | None = None
    notes: list[str] = field(default_factory=list)
    cards: list["Block"] = field(default_factory=list)


def parse_source(path: Path) -> tuple[list[Block], list[str]]:
    sections: list[Block] = []
    refs: list[str] = []
    section = card = None
    in_refs = False
    para: list[str] = []

    def flush() -> None:
        if para and section is not None:
            (card or section).notes.append(" ".join(para))
        para.clear()

    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("## "):
            flush()
            title = re.sub(r"^\d+\.\s*", "", line[3:]).strip()
            in_refs = title == "References"
            section = None if in_refs else Block(title)
            card = None
            if section:
                sections.append(section)
            continue
        if in_refs:
            if line.startswith("- "):
                refs.append(line[2:].strip())
            continue
        if section is None:
            continue
        target = card or section
        if line.startswith("### "):
            flush()
            card = Block(line[4:].strip())
            section.cards.append(card)
        elif line.startswith("> "):
            section.takeaway = line[2:].strip()
        elif line.startswith("$$"):
            target.eqs.append(line.strip().strip("$").strip())
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                target.table = (target.table or []) + [cells]
        elif line.startswith("- "):
            target.bullets.append(line[2:].strip())
        elif re.match(r"\s+- ", line):
            target.notes.append(line.strip()[2:].strip())
        elif not line.strip():
            flush()
        else:
            para.append(line.strip())
    flush()
    return sections, refs


def notes_for(section: Block) -> str:
    parts = [section.takeaway] if section.takeaway else []
    parts += section.notes
    for c in section.cards:
        if c.notes:
            parts.append(f"{c.title}: " + " ".join(c.notes))
    return "\n\n".join(parts)


# --------------------------------------------------------------------------- drawing helpers


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def runs_from_markup(text: str, bold_label: bool) -> list[tuple[str, bool, bool]]:
    """Split **bold** / *italic* markup; optionally bold a short 'Label:' prefix."""
    out: list[tuple[str, bool, bool]] = []
    m = re.match(r"^([^:*]{2,28}): (.+)$", text) if bold_label else None
    if m:
        out.append((m.group(1) + ": ", True, False))
        text = m.group(2)
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if tok.startswith("**"):
            out.append((tok[2:-2], True, False))
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 1:
            out.append((tok[1:-1], False, True))
        elif tok:
            out.append((tok, False, False))
    return out


def set_bullet(p, color: str = BLUE, indent: float = 0.15) -> None:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    bu_clr = pPr.makeelement(qn("a:buClr"), {})
    bu_clr.append(bu_clr.makeelement(qn("a:srgbClr"), {"val": color}))
    pPr.append(bu_clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def add_par(tf, text, size, color=INK, bold=False, italic=False, font=BODY_FONT,
            bullet=False, align=None, after=3, bold_label=False):
    fresh = len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs
    p = tf.paragraphs[0] if fresh else tf.add_paragraph()
    p.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    for chunk, b, i in runs_from_markup(text, bold_label):
        r = p.add_run()
        r.text = chunk
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold or b, italic or i, font
        f.color.rgb = rgb(color)
    if bullet:
        set_bullet(p)
    return p


def textbox(slide, x, y, w, h, margin=0.04, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    return tf


def shape(slide, x, y, w, h, fill, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, line=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    style = s._element.find(qn("p:style"))  # drop the theme's shadow/effect style
    if style is not None:
        s._element.remove(style)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def shape_text(s, lines, anchor=MSO_ANCHOR.MIDDLE, margin=0.08):
    """lines: list of (text, size, color, bold, italic, align)."""
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for side in ("top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.03))
    for text, size, color, bold, italic, align in lines:
        add_par(tf, text, size, color=color, bold=bold, italic=italic, align=align, after=1)


def bullets(tf, items, size, bold_label=True, after=4):
    for item in items:
        add_par(tf, item, size, bullet=True, after=after, bold_label=bold_label)


_eq_cache: dict[tuple[str, float], tuple[Path, float, float]] = {}


def eq_png(latex: str, fs: float) -> tuple[Path, float, float]:
    key = (latex, fs)
    if key not in _eq_cache:
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex}$", fontsize=fs, color="#" + INK)
        path = TMP / f"eq{len(_eq_cache)}.png"
        fig.savefig(path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.03)
        plt.close(fig)
        w, h = Image.open(path).size
        _eq_cache[key] = (path, w / 300, h / 300)
    return _eq_cache[key]


def place_eq(slide, latex, x, y, w, h, fs=14.0):
    """Centre an equation in the box; stack its \\qquad parts if it would shrink too much."""
    parts = [latex]
    _, nw, nh = eq_png(latex, fs)
    if min(w / nw, h / nh) < 0.8 and r"\qquad" in latex:
        parts = [p.strip().rstrip(",") for p in latex.split(r"\qquad")]
    imgs = [eq_png(p, fs) for p in parts]
    gap = 0.04
    tot_h = sum(i[2] for i in imgs) + gap * (len(imgs) - 1)
    scale = min(1.0, w / max(i[1] for i in imgs), h / tot_h)
    cy = y + (h - tot_h * scale) / 2
    for path, iw, ih in imgs:
        slide.shapes.add_picture(str(path), Inches(x + (w - iw * scale) / 2), Inches(cy),
                                 Inches(iw * scale), Inches(ih * scale))
        cy += (ih + gap) * scale


def table(slide, x, y, rows, col_w, size=10.5, row_h=0.3, highlight=None):
    n, m = len(rows), len(rows[0])
    frame = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * n))
    tbl = frame.table
    tbl._tbl.tblPr.set("bandRow", "0")
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        hl = highlight is not None and i > 0 and highlight(row)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(NAVY if i == 0 else GREEN_TINT if hl else TINT if i % 2 else WHITE)
            tf = c.text_frame
            tf.word_wrap = True
            if i == 0:
                add_par(tf, val, size, color=WHITE, bold=True, after=0)
            else:
                add_par(tf, val, size, color=INK, bold=(j == 0), after=0)
    return frame


def card(slide, x, y, w, h, title, items=(), size=11.0, eqs=(), eq_h=0.7, eq_fs=13.0,
         header=NAVY, body=TINT):
    """Navy pill header over a tinted body: bullets on top, equations pinned to the bottom."""
    shape_text(shape(slide, x, y, w, 0.34, header, radius=0.25),
               [(title, 12, WHITE, True, False, PP_ALIGN.LEFT)], margin=0.12)
    shape(slide, x, y + 0.4, w, h - 0.4, body, radius=0.04)
    eq_total = eq_h * len(eqs)
    if items:
        tf = textbox(slide, x + 0.08, y + 0.46, w - 0.16, h - 0.5 - eq_total)
        bullets(tf, items, size)
    for k, latex in enumerate(eqs):
        place_eq(slide, latex, x + 0.1, y + h - eq_total + k * eq_h, w - 0.2, eq_h - 0.08, fs=eq_fs)


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(10), Inches(7.5)
        self.n = 0

    def slide(self, title, takeaway="", owner=None, notes=""):
        self.n += 1
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        add_par(textbox(s, 0.45, 0.28, 9.1, 0.55, anchor=MSO_ANCHOR.MIDDLE), title, 24,
                color=BLUE, bold=True, font=TITLE_FONT, after=0)
        if takeaway:
            add_par(textbox(s, 0.45, 0.86, 9.1, 0.55), takeaway, 12.5, italic=True, after=0)
        add_par(textbox(s, 0.45, 7.1, 6.0, 0.25), FOOTER, 8.5, color=MUTED, after=0)
        right = f"{owner}  ·  {self.n}" if owner else str(self.n)
        add_par(textbox(s, 6.55, 7.1, 3.0, 0.25), right, 8.5, color=MUTED, after=0, align=PP_ALIGN.RIGHT)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s


def paper_figure(page: int) -> Path | None:
    """Extract a figure image from the paper PDF and trim its white margins."""
    if not shutil.which("pdfimages") or not PAPER_PDF.exists():
        return None
    prefix = TMP / f"fig{page}"
    subprocess.run(["pdfimages", "-f", str(page), "-l", str(page), "-j", str(PAPER_PDF), str(prefix)],
                   check=True)
    found = sorted(TMP.glob(f"fig{page}-*"))
    if not found:
        return None
    img = Image.open(found[0]).convert("RGB")
    bbox = ImageChops.difference(img, Image.new("RGB", img.size, "white")).getbbox()
    out = TMP / f"figure_p{page}.png"
    (img.crop(bbox) if bbox else img).save(out)
    return out


# --------------------------------------------------------------------------- team sections


def team_intro(d: Deck) -> None:
    s = d.slide("Topic & Brief Introduction")
    tf = textbox(s, 0.45, 1.0, 5.2, 5.9)
    add_par(tf, "Topic:", 14, bold=True, after=2)
    add_par(tf, "Extending “Pairs Trading: Performance of a Relative-Value Arbitrage Rule” by "
                "Gatev, Goetzmann & Rouwenhorst, Review of Financial Studies (2006).", 13, after=12)
    add_par(tf, "Brief introduction:", 14, bold=True, after=2)
    bullets(tf, [
        "Tests a Wall Street rule: pair stocks whose normalized prices moved together, "
        "then bet on convergence when they diverge",
        "U.S. daily data 1962–2002: top pairs earn up to ~11% a year in excess returns, "
        "above conservative cost estimates",
        "Not explained by standard risk factors or simple reversal, but returns fall after the late 1980s",
        "Our project: replicate the core experiment, then extend it to 2026 with explicit "
        "trading frictions, stress tests and new pair-selection methods",
    ], 13, bold_label=False, after=6)
    add_par(tf, "", 6)
    add_par(tf, "Team: Jared Cohen, Jacob Waxman, Eiko Reisz, Hao Wang", 13, bold_label=True, after=0)
    fig = paper_figure(44)
    if fig:
        w = 3.75
        iw, ih = Image.open(fig).size
        s.shapes.add_picture(str(fig), Inches(5.8), Inches(1.6), Inches(w), Inches(w * ih / iw))
        add_par(textbox(s, 5.8, 1.7 + w * ih / iw, w, 0.9),
                "Figure 1 (GGR 2006): Kennecott vs. Uniroyal normalized prices, Aug 1963–Jan 1964. "
                "The pair opens and closes several times in one trading period.", 9.5, color=MUTED, italic=True)


def team_summary(d: Deck) -> None:
    s = d.slide("Summary of Paper", "Do temporary divergences between historically similar stocks "
                "create relative-value arbitrage profits?")
    tf = textbox(s, 0.45, 1.5, 4.45, 5.4)
    bullets(tf, [
        "Idea: a broader Law of One Price; near-substitutes that drift apart should converge again",
        "Method: form pairs over 12 months by minimum distance between normalized total-return "
        "prices; trade them over the next 6 months",
        "Rule: buy the loser and short the winner at a 2σ divergence; close when the prices cross",
        "Finding: profits survive conservative cost estimates and are not explained by known "
        "factors or simple reversal",
        "Caveat: profitability declines in the later years of the sample",
    ], 13, after=8)
    data = CategoryChartData()
    data.categories = ["Top 5", "Top 20", "Pairs 101–120", "All pairs"]
    data.add_series("Trade on signal day", (1.31, 1.44, 1.08, 1.10))
    data.add_series("Wait one day", (0.75, 0.90, 0.80, 0.72))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.05), Inches(1.5),
                            Inches(4.5), Inches(4.6), data)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = "Mean monthly excess return (%)"
    title_font = ch.chart_title.text_frame.paragraphs[0].runs[0].font
    title_font.size, title_font.bold, title_font.name = Pt(12), True, BODY_FONT
    title_font.color.rgb = rgb(INK)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.name = Pt(10.5), BODY_FONT
    for series, color in zip(ch.plots[0].series, (NAVY, "8DB3DD")):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(color)
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "0.00", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.name = Pt(10), BODY_FONT
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.tick_labels.font.size = Pt(10.5)
    ch.category_axis.tick_labels.font.name = BODY_FONT
    ch.category_axis.format.line.color.rgb = rgb("B8C4D2")
    add_par(textbox(s, 5.05, 6.15, 4.5, 0.75),
            "Source: GGR (2006) Table 1, fully invested, Jul 1963–Dec 2002. Waiting one day removes "
            "most of the bid-ask bounce; profits stay positive and significant.", 9.5, color=MUTED, italic=True)


def team_data(d: Deck) -> None:
    s = d.slide("Data Sourcing", "Daily CRSP data for all U.S. common stocks, including firms that "
                "later delist; factors and the T-bill rate from the Ken French Data Library.")
    w, h = 4.475, 2.68
    card(s, 0.45, 1.5, w, h, "Sources & sample", [
        "CRSP daily (WRDS): share codes 10–11 on NYSE, AMEX and NASDAQ, delisted firms included",
        "Replication: Jul 1963–Dec 2002; extension through the latest CRSP release (Mar 2026)",
        "Backups if CRSP access is limited: FMP, Alpha Vantage, Massive",
    ], size=12)
    card(s, 5.075, 1.5, w, h, "Variables", [
        "Primary: RET (with dividends), DLRET, price, volume, shares outstanding, SIC, PERMNO",
        "Constructed: total-return index; market cap → NYSE size deciles; SIC → sectors",
    ], size=12, eqs=[r"P_{i,t}=\prod_{s=t_0+1}^{t}(1+r_{i,s}),\quad P_{i,0}=1"], eq_h=0.62)
    card(s, 0.45, 4.3, w, h, "Filtering & missing data", [
        "Exclude stocks with any zero-volume or missing-return day in the formation window",
        "Delisted positions close at DLRET; −30% if it is missing for performance delistings",
        "Missing returns while listed are set to zero; nothing is interpolated",
    ], size=12)
    card(s, 5.075, 4.3, w, h, "Limitations", [
        "Daily closes ignore intraday execution; costs are inferred by waiting one day",
        "Borrowing costs and short recalls are not observed",
        "The ~23-year extension is shorter than the paper's 40 years and dominated by a few crises",
        "Fallback vendors may under-cover delisted firms (survivorship bias)",
    ], size=12)


def team_pairs(d: Deck) -> None:
    s = d.slide("Pairs Construction Methodology", "Start from the paper's mechanical rules, then "
                "make pair selection adaptive.", owner="Jared Cohen (draft)")
    card(s, 0.45, 1.5, 4.475, 4.6, "Paper's rules (baseline)", [
        "Entry: open a long–short position when normalized prices diverge by more than "
        "2 formation-period standard deviations",
        "Convergence exit: close when the two price paths cross again",
        "Time exit: if they never cross, close at the end of the 6-month trading period",
        "Delisting exit: close at the delisting return or the last available price",
        "Risk analysis: returns, worst losses, risk factors, transaction costs and short-selling "
        "risks; do not lean on historical VaR",
    ], size=14)
    card(s, 5.075, 1.5, 4.475, 4.6, "Our ideas (in progress)", [
        "Build pairs on each stock's total-return index, dividends included",
        "ML-based pair selection (specified on slide 12)",
        "Regime-based selection: choose dynamically which pairs to trade",
        "Structural-break test with risk factors: drop pairs that stop moving together and "
        "close their positions",
    ], size=14, header=GREEN, body=GREEN_TINT)


def team_production(d: Deck) -> None:
    s = d.slide("Production Deployment & Data Architecture", "A modular, reproducible research "
                "pipeline on Google Cloud Platform.")
    steps = [("Ingest", "CRSP daily, security metadata, French factors"),
             ("Validate", "Missing data, corporate actions, duplicates"),
             ("Store", "GCS Parquet: raw, cleaned, derived"),
             ("Query", "BigQuery: checks, analysis, reports"),
             ("Backtest", "Python strategy on Cloud Run Jobs")]
    bw, gap = 1.62, 0.25
    for k, (head, body) in enumerate(steps):
        x = 0.45 + k * (bw + gap)
        shape_text(shape(s, x, 1.55, bw, 1.15, NAVY if k < 4 else GREEN),
                   [(head, 13, WHITE, True, False, PP_ALIGN.CENTER),
                    (body, 10, WHITE, False, False, PP_ALIGN.CENTER)])
        if k < len(steps) - 1:
            shape(s, x + bw + 0.04, 1.98, gap - 0.08, 0.3, BLUE, kind=MSO_SHAPE.RIGHT_ARROW)
    card(s, 0.45, 3.05, 4.475, 3.9, "Research pipeline", [
        "Raw, cleaned and derived datasets stay separate",
        "Versioned configs and timestamped outputs: reproducible, no look-ahead",
        "Python handles pair formation, signals, portfolio construction and backtesting",
        "Cloud Run Jobs schedule ingestion and backtest runs",
    ], size=13)
    card(s, 5.075, 3.05, 4.475, 3.9, "Production extension", [
        "Refresh the eligible universe and re-rank pairs on a schedule",
        "Monitor open positions and portfolio drawdowns",
        "Alerts and predefined exits when spreads or risk exposures breach thresholds",
        "Initial scope is historical research, not live execution",
    ], size=13)


def team_feasibility(d: Deck) -> None:
    s = d.slide("Project Feasibility & Possible Hurdles", "Feasible with data already in the repo: "
                "CRSP daily 1961–Mar 2026 and French factors are acquired and validated.")
    rows = [["Hurdle", "How we address it"],
            ["Transaction costs and short-selling constraints", "Model trading and financing costs explicitly (scenarios S1–S3)"],
            ["Exposure to systematic risk factors", "Report factor-adjusted alpha; the factor gate for every extension"],
            ["Profitability varies over time", "Compare subperiods; extend the sample to 2026"],
            ["Look-ahead and survivorship bias", "Point-in-time CRSP with delisted firms; formation-window data only"],
            ["Regime changes; pairs that never re-converge", "Regime-aware exits, stress tests, out-of-sample testing"],
            ["Overfitting from repeated testing", "Predefined parameters, locked test window, sensitivity analysis"],
            ["Data vintage differs from the 2006 database", "Document differences; report raw and risk-adjusted returns"]]
    table(s, 0.45, 1.55, rows, [3.6, 5.5], size=12, row_h=0.62)


def references(d: Deck, refs: list[str]) -> None:
    s = d.slide("References")
    tf = textbox(s, 0.45, 1.0, 9.1, 6.0)
    for ref in refs:
        add_par(tf, ref, 10.5, after=5)


# --------------------------------------------------------------------------- Hao's sections


def hao_slides(d: Deck, secs: list[Block]) -> None:
    expected = ["Research Questions", "Core Experiment", "Microstructure Frictions", "Explaining the Returns",
                "Out-of-Sample Design", "Risk Metrics", "Proposed Extensions", "Machine-Learning"]
    assert len(secs) == len(expected), [s.title for s in secs]
    for sec, prefix in zip(secs, expected):
        assert sec.title.startswith(prefix), (sec.title, prefix)
    q_map, sim, fric, factors, oos, risk, ext, ml = secs

    def new(sec: Block):
        return d.slide(sec.title, sec.takeaway, owner=HAO, notes=notes_for(sec))

    # Research questions -> experiments
    s = new(q_map)
    table(s, 0.45, 1.5, q_map.table, [2.0, 2.65, 0.7, 3.75], size=10.5, row_h=0.5,
          highlight=lambda r: "gate" in r[0])
    for k, text in enumerate(q_map.bullets):
        gate = "gate" in text
        box_ = shape(s, 0.45 + k * 4.625, 5.85, 4.475, 0.95, GREEN_TINT if gate else TINT2, radius=0.12)
        box_.text_frame.word_wrap = True
        tf = box_.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.12)
        add_par(tf, text, 12, after=0, bold_label=True, align=PP_ALIGN.LEFT)

    # Core simulation
    s = new(sim)
    w, gap = 2.93, 0.155
    for k, c in enumerate(sim.cards[:3]):
        card(s, 0.45 + k * (w + gap), 1.5, w, 3.9, c.title, c.bullets, size=11.5, eqs=c.eqs,
             eq_h=0.95, eq_fs=12.5)
    io = sim.cards[3]
    card(s, 0.45, 5.52, 9.1, 1.46, io.title, io.bullets, size=11)

    # Frictions: scenario ladder + two cards
    s = new(fric)
    bw, aw = 2.05, 0.3
    for k, (name, what, measures) in enumerate(fric.table[1:]):
        x = 0.45 + k * (bw + aw)
        shape_text(shape(s, x, 1.5, bw, 1.12, GREEN if k == 3 else NAVY),
                   [(name, 13, WHITE, True, False, PP_ALIGN.LEFT),
                    (what, 10, WHITE, False, False, PP_ALIGN.LEFT),
                    (f"→ {measures}", 10, "CFE0F2" if k < 3 else "D5F0E4", False, True, PP_ALIGN.LEFT)],
                   anchor=MSO_ANCHOR.TOP)
        if k < 3:
            shape(s, x + bw + 0.05, 1.92, aw - 0.1, 0.28, BLUE, kind=MSO_SHAPE.RIGHT_ARROW)
    rep, extn = fric.cards
    card(s, 0.45, 2.8, 3.85, 4.17, rep.title, rep.bullets, size=12.5, eqs=rep.eqs, eq_h=0.8)
    card(s, 4.45, 2.8, 5.1, 4.17, extn.title, extn.bullets, size=12.5, eqs=extn.eqs, eq_h=0.75, eq_fs=12.5)

    # Factors / reversal / decline
    s = new(factors)
    shape(s, 0.45, 1.5, 9.1, 0.66, TINT2, radius=0.12)
    place_eq(s, factors.eqs[0], 0.6, 1.52, 8.8, 0.62, fs=14)
    for k, c in enumerate(factors.cards):
        gate = c.title.startswith("Extend")
        card(s, 0.45 + k * (w + gap), 2.35, w, 4.62, c.title, c.bullets, size=12.5,
             header=GREEN if gate else NAVY, body=GREEN_TINT if gate else TINT)

    # Out-of-sample design: timeline + cards
    s = new(oos)
    seg_w = [3.9, 2.2, 3.0]
    fills = [TINT2, BLUE, NAVY]
    x = 0.45
    add_par(textbox(s, 0.45, 1.45, sum(seg_w[:2]), 0.26), "Paper's sample: replication, 1963–2002",
            10, color=MUTED, italic=True, after=0, align=PP_ALIGN.CENTER)
    add_par(textbox(s, 0.45 + sum(seg_w[:2]), 1.45, seg_w[2], 0.26), "New data: extension, 2003–2026",
            10, color=MUTED, italic=True, after=0, align=PP_ALIGN.CENTER)
    for k, (window, period, use) in enumerate(oos.table[1:]):
        dark = k > 0
        fg = WHITE if dark else INK
        shape_text(shape(s, x, 1.75, seg_w[k] - 0.05, 0.55, fills[k], radius=0.1),
                   [(window, 11.5, fg, True, False, PP_ALIGN.CENTER),
                    (period, 10.5, fg, False, False, PP_ALIGN.CENTER)])
        add_par(textbox(s, x, 2.34, seg_w[k] - 0.05, 0.62), use, 10, after=0, align=PP_ALIGN.CENTER)
        x += seg_w[k]
    for k, c in enumerate(oos.cards):
        card(s, 0.45 + k * (w + gap), 3.05, w, 3.92, c.title, c.bullets, size=13)

    # Risk & stress
    s = new(risk)
    metrics, stress, shocks = risk.cards
    card(s, 0.45, 1.5, 4.2, 2.6, metrics.title, metrics.bullets, size=11)
    card(s, 0.45, 4.25, 4.2, 2.72, shocks.title, shocks.bullets, size=11, eqs=shocks.eqs, eq_h=0.6)
    card(s, 4.8, 1.5, 4.75, 5.47, stress.title)
    table(s, 4.88, 1.98, stress.table, [1.38, 1.5, 1.71], size=10, row_h=0.42)
    tf = textbox(s, 4.9, 5.55, 4.55, 1.3)
    bullets(tf, stress.bullets, 11)

    # Extensions
    s = new(ext)
    table(s, 0.45, 1.5, ext.table, [0.45, 3.3, 2.75, 2.6], size=10.5, row_h=0.55)
    c = ext.cards[0]
    card(s, 0.45, 5.0, 9.1, 1.98, c.title, c.bullets, size=12)

    # ML specification: 2x2 grid
    s = new(ml)
    w2 = 4.475
    for k, c in enumerate(ml.cards):
        x, y = 0.45 + (k % 2) * (w2 + 0.15), 1.5 + (k // 2) * 2.78
        card(s, x, y, w2, 2.66, c.title, c.bullets, size=12.5, eqs=c.eqs, eq_h=0.55, eq_fs=12.5)


def main() -> None:
    secs, refs = parse_source(SOURCE_MD)
    d = Deck()
    team_intro(d)
    team_summary(d)
    team_data(d)
    team_pairs(d)
    hao_slides(d, secs)
    team_production(d)
    team_feasibility(d)
    references(d, refs)
    d.prs.save(OUT)
    shutil.rmtree(TMP, ignore_errors=True)
    print(f"wrote {OUT.relative_to(ROOT)} ({d.n} slides)")


if __name__ == "__main__":
    main()
